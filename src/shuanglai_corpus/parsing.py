"""Offline-friendly text and metadata extraction."""

from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin, urlparse

DATE_PATTERNS = [
    re.compile(r"(?P<y>20\d{2})[年./-](?P<m>\d{1,2})[月./-](?P<d>\d{1,2})日?"),
    re.compile(r"(?P<y>20\d{2})年(?P<m>\d{1,2})月"),
]
DOC_NUMBER_RE = re.compile(r"[\u4e00-\u9fff]{1,8}(?:字|发|政发|办发|函|批|复|令)?[〔\[]20\d{2}[〕\]]\s*\d+号")
AUTHORITY_RE = re.compile(r"(?:发布机构|发文机关|来源)\s*[：:]\s*([^\n|]{2,40})")
ATTACHMENT_EXTENSIONS = (".pdf", ".doc", ".docx", ".xls", ".xlsx", ".txt", ".wps")


class _FallbackHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.text_parts, self.title_parts, self.h1_parts, self.links = [], [], [], []
        self.in_title = self.in_h1 = False
        self.suppressed = 0

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in ("script", "style", "noscript", "nav", "footer"): self.suppressed += 1
        if tag == "title": self.in_title = True
        if tag == "h1": self.in_h1 = True
        if tag == "a" and attrs.get("href"): self.links.append((attrs["href"], ""))

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript", "nav", "footer") and self.suppressed: self.suppressed -= 1
        if tag == "title": self.in_title = False
        if tag == "h1": self.in_h1 = False

    def handle_data(self, data):
        value = data.strip()
        if not value or self.suppressed: return
        self.text_parts.append(value)
        if self.in_title: self.title_parts.append(value)
        if self.in_h1: self.h1_parts.append(value)


def parse_chinese_date(text: str) -> str:
    for pattern in DATE_PATTERNS:
        match = pattern.search(text or "")
        if match:
            year, month = int(match["y"]), int(match["m"])
            day = int(match.groupdict().get("d") or 1)
            return f"{year:04d}-{month:02d}-{day:02d}"
    return ""


def extract_document_number(text: str) -> str:
    match = DOC_NUMBER_RE.search(text or "")
    return re.sub(r"\s+", "", match.group(0)) if match else ""


def extract_issuing_authority(text: str, title: str = "") -> str:
    match = AUTHORITY_RE.search(text or "")
    if match:
        return match.group(1).strip()
    for name in ("山东省人民政府", "青岛市人民政府", "烟台市人民政府", "莱西市人民政府", "莱阳市人民政府", "发展和改革委员会"):
        if name in (title + " " + text[:1000]):
            return name
    return ""


def extract_html(html: bytes | str, base_url: str = "") -> tuple[str, str, list[dict]]:
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        parser = _FallbackHTMLParser()
        parser.feed(html.decode("utf-8", errors="replace") if isinstance(html, bytes) else html)
        title = " ".join(parser.h1_parts or parser.title_parts)
        attachments = []
        for href, label in parser.links:
            absolute = urljoin(base_url, href)
            if urlparse(absolute).path.lower().endswith(ATTACHMENT_EXTENSIONS):
                attachments.append({"url": absolute, "filename": Path(urlparse(absolute).path).name, "label": label})
        return title, "\n".join(parser.text_parts), attachments

    soup = BeautifulSoup(html, "lxml")
    for element in soup(["script", "style", "noscript", "nav", "footer"]):
        element.decompose()
    title = ""
    heading = soup.find("h1")
    if heading:
        title = heading.get_text(" ", strip=True)
    elif soup.title:
        title = soup.title.get_text(" ", strip=True)
    selectors = ["article", ".article", ".content", ".TRS_Editor", "#zoom", "main"]
    container = next((soup.select_one(selector) for selector in selectors if soup.select_one(selector)), soup.body or soup)
    text = container.get_text("\n", strip=True)
    attachments = []
    for anchor in soup.find_all("a", href=True):
        href = urljoin(base_url, anchor["href"])
        path = urlparse(href).path.lower()
        if path.endswith(ATTACHMENT_EXTENSIONS):
            attachments.append({"url": href, "filename": Path(urlparse(href).path).name, "label": anchor.get_text(" ", strip=True)})
    return title, text, attachments


def extract_file(path: Path) -> tuple[str, str]:
    suffix = path.suffix.lower()
    if suffix in (".html", ".htm"):
        _, text, _ = extract_html(path.read_bytes())
        return text, "parsed"
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace"), "parsed"
    if suffix == ".pdf":
        from pypdf import PdfReader
        text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages).strip()
        return (text, "parsed") if text else ("", "needs_ocr")
    if suffix == ".docx":
        from docx import Document
        return "\n".join(p.text for p in Document(str(path)).paragraphs), "parsed"
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        text = "\n".join("\t".join("" if value is None else str(value) for value in row) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True))
        return text, "parsed"
    return "", "unsupported_format"
